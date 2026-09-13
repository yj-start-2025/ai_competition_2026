import io
import json
import re
import time
from pathlib import Path

import pandas as pd
import streamlit as st
from openai import OpenAI

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont


# =========================================================
# 0. 기본 설정
# =========================================================
st.set_page_config(
    page_title="수강후기 AI 분석 리포트",
    page_icon="📊",
    layout="wide",
)

APP_TITLE = "수강후기 AI 분석 리포트"
MODEL = "gpt-5.6-luna"
SAMPLE_FILE = Path(__file__).with_name("reviews_seoul_center_final.csv")

st.title(f"📊 {APP_TITLE}")
st.caption(
    "수강후기 데이터를 기반으로 과정별 핵심 키워드, 감성, 주요 의견, "
    "개선방안 및 개인별 피드백을 자동으로 분석합니다."
)


# =========================================================
# 1. OpenAI API
# =========================================================
try:
    client = OpenAI(api_key=st.secrets["OPENAI_API_KEY"])
except Exception:
    st.error(
        "OpenAI API 키가 설정되지 않았습니다. "
        "Streamlit Cloud → App settings → Secrets에 "
        'OPENAI_API_KEY = "sk-..." 형식으로 등록해주세요.'
    )
    st.stop()


# =========================================================
# 2. 데이터 유틸리티
# =========================================================
def normalize_name(value):
    return (
        str(value)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )


def find_col(df, exact_candidates, partial_candidates=None, exclude=None):
    partial_candidates = partial_candidates or []
    exclude = {normalize_name(x) for x in (exclude or [])}
    normalized = {col: normalize_name(col) for col in df.columns}

    for candidate in exact_candidates:
        c = normalize_name(candidate)
        for col, ncol in normalized.items():
            if ncol in exclude:
                continue
            if ncol == c:
                return col

    for candidate in partial_candidates:
        c = normalize_name(candidate)
        for col, ncol in normalized.items():
            if ncol in exclude:
                continue
            if c in ncol or ncol in c:
                return col

    return None


def load_dataframe(file_obj):
    name = file_obj.name.lower()

    if name.endswith(".csv"):
        try:
            return pd.read_csv(file_obj)
        except UnicodeDecodeError:
            file_obj.seek(0)
            return pd.read_csv(file_obj, encoding="cp949")

    return pd.read_excel(file_obj)


@st.cache_data(show_spinner=False)
def load_sample(path_str):
    return pd.read_csv(path_str)


def clean_dataframe(df, course_col, review_col, rating_col):
    work = df.copy()
    original_n = len(work)

    work = work.dropna(subset=[course_col, review_col])
    work[course_col] = work[course_col].astype(str).str.strip()
    work[review_col] = work[review_col].astype(str).str.strip()

    # 의미 없는 공백만 제거하고, 짧은 후기 자체는 보존
    work = work[work[review_col].str.len() >= 1]

    if rating_col:
        work[rating_col] = pd.to_numeric(work[rating_col], errors="coerce")

    removed_n = original_n - len(work)
    return work.reset_index(drop=True), removed_n


def extract_json(text):
    if not text:
        raise ValueError("모델 응답이 비어 있습니다.")

    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    start = text.find("{")
    end = text.rfind("}")

    if start == -1 or end == -1 or end <= start:
        raise ValueError("모델 응답에서 JSON 객체를 찾지 못했습니다.")

    return json.loads(text[start:end + 1])


def safe_list(obj, key, limit=None):
    value = obj.get(key, [])
    if not isinstance(value, list):
        return []
    return value[:limit] if limit else value


def compact_text(text):
    return re.sub(r"\s+", "", str(text).lower())


def compute_keyword_frequencies(sample, review_col, keywords):
    review_texts = [compact_text(x) for x in sample[review_col].astype(str)]

    for kw in keywords:
        terms = [kw.get("word", "")]
        aliases = kw.get("aliases", [])

        if isinstance(aliases, list):
            terms.extend(aliases)

        terms = [compact_text(x) for x in terms if str(x).strip()]
        terms = list(dict.fromkeys(terms))

        count = 0
        for text in review_texts:
            if any(term and term in text for term in terms):
                count += 1

        kw["frequency"] = count

    return keywords


def build_course_overview(df, course_col, rating_col):
    if rating_col and df[rating_col].notna().any():
        overview = (
            df.groupby(course_col)
            .agg(
                후기수=(course_col, "size"),
                평균별점=(rating_col, "mean"),
            )
            .reset_index()
            .rename(columns={course_col: "과정명"})
        )
        overview["평균별점"] = overview["평균별점"].round(2)
    else:
        overview = (
            df.groupby(course_col)
            .size()
            .reset_index(name="후기수")
            .rename(columns={course_col: "과정명"})
        )
        overview["평균별점"] = None

    return overview.sort_values("후기수", ascending=False).reset_index(drop=True)


# =========================================================
# 3. 리포트 생성
# =========================================================
def result_to_rows(result):
    counts = result.get("sentiment_counts", {})
    return {
        "positive": int(counts.get("positive", 0) or 0),
        "neutral": int(counts.get("neutral", 0) or 0),
        "negative": int(counts.get("negative", 0) or 0),
        "analysis_n": int(result.get("analysis_review_count", 0) or 0),
    }


def build_docx_report(results):
    doc = Document()

    styles = doc.styles
    styles["Normal"].font.name = "Malgun Gothic"
    styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("수강후기 AI 분석 리포트")
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph(
        "생성형 AI를 활용하여 과정별 핵심 키워드, 감성, 주요 의견, "
        "개선방안 및 개인별 피드백을 자동 분석한 결과입니다."
    )

    for idx, (course, result) in enumerate(results.items(), start=1):
        doc.add_heading(f"{idx}. {course}", level=1)

        summary = result.get("executive_summary", "")
        if summary:
            doc.add_heading("AI 핵심 요약", level=2)
            doc.add_paragraph(summary)

        counts = result_to_rows(result)
        doc.add_heading("감성 분석", level=2)

        table = doc.add_table(rows=2, cols=4)
        table.style = "Table Grid"
        headers = ["분석 후기", "긍정", "중립", "부정"]
        values = [
            counts["analysis_n"],
            counts["positive"],
            counts["neutral"],
            counts["negative"],
        ]
        for j, h in enumerate(headers):
            table.cell(0, j).text = h
            table.cell(1, j).text = str(values[j])

        doc.add_heading("핵심 키워드 TOP 10", level=2)
        kw_table = doc.add_table(rows=1, cols=3)
        kw_table.style = "Table Grid"
        kw_table.rows[0].cells[0].text = "키워드"
        kw_table.rows[0].cells[1].text = "빈도"
        kw_table.rows[0].cells[2].text = "의미"

        for kw in result.get("keywords", [])[:10]:
            cells = kw_table.add_row().cells
            cells[0].text = str(kw.get("word", ""))
            cells[1].text = str(kw.get("frequency", 0))
            cells[2].text = str(kw.get("meaning", ""))

        reasons = result.get("sentiment_reasons", {})
        for label, key in [
            ("긍정 주요 사유", "positive"),
            ("중립 주요 사유", "neutral"),
            ("부정·개선 요구 주요 사유", "negative"),
        ]:
            doc.add_heading(label, level=2)
            vals = safe_list(reasons, key, 5)
            if vals:
                for v in vals:
                    doc.add_paragraph(str(v), style="List Bullet")
            else:
                doc.add_paragraph("해당 의견 없음")

        doc.add_heading("개선 방안 5가지", level=2)
        for i, item in enumerate(
            result.get("improvement_suggestions", [])[:5],
            start=1,
        ):
            doc.add_paragraph(f"{i}. {item}")

        doc.add_heading("개인별 피드백", level=2)
        for fb in result.get("individual_feedback", [])[:10]:
            rid = fb.get("review_id", "")
            snippet = fb.get("review_snippet", "")
            feedback = fb.get("feedback", "")
            doc.add_paragraph(
                f"교육생 의견 #{rid} - {snippet}",
                style="List Bullet",
            )
            doc.add_paragraph(f"AI 피드백: {feedback}")

        if idx < len(results):
            doc.add_page_break()

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(
        "본 결과는 생성형 AI 기반 자동 분석 결과이며, 실제 업무 적용 시 담당자 최종 검토를 권장합니다."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_pdf_report(results):
    # ReportLab 기본 CID Korean font
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
        pdfmetrics.registerFont(UnicodeCIDFont("HYGoThic-Medium"))
        body_font = "HYSMyeongJo-Medium"
        bold_font = "HYGoThic-Medium"
    except Exception:
        body_font = "Helvetica"
        bold_font = "Helvetica-Bold"

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "KTitle",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=18,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=10,
    )

    h1 = ParagraphStyle(
        "KH1",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=14,
        leading=19,
        spaceBefore=8,
        spaceAfter=6,
    )

    h2 = ParagraphStyle(
        "KH2",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=11,
        leading=15,
        spaceBefore=6,
        spaceAfter=4,
    )

    body = ParagraphStyle(
        "KBody",
        parent=styles["BodyText"],
        fontName=body_font,
        fontSize=9,
        leading=13,
        spaceAfter=4,
    )

    story = [
        Paragraph("수강후기 AI 분석 리포트", title_style),
        Paragraph(
            "생성형 AI를 활용하여 과정별 핵심 키워드, 감성, 주요 의견, "
            "개선방안 및 개인별 피드백을 자동 분석한 결과입니다.",
            body,
        ),
        Spacer(1, 4 * mm),
    ]

    for idx, (course, result) in enumerate(results.items(), start=1):
        story.append(Paragraph(f"{idx}. {course}", h1))

        summary = result.get("executive_summary", "")
        if summary:
            story.append(Paragraph("AI 핵심 요약", h2))
            story.append(Paragraph(summary, body))

        counts = result_to_rows(result)

        story.append(Paragraph("감성 분석", h2))
        sentiment_data = [
            ["분석 후기", "긍정", "중립", "부정"],
            [
                str(counts["analysis_n"]),
                str(counts["positive"]),
                str(counts["neutral"]),
                str(counts["negative"]),
            ],
        ]
        sentiment_table = Table(
            sentiment_data,
            colWidths=[35 * mm, 35 * mm, 35 * mm, 35 * mm],
        )
        sentiment_table.setStyle(
            TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), body_font),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ])
        )
        story.append(sentiment_table)
        story.append(Spacer(1, 3 * mm))

        story.append(Paragraph("핵심 키워드 TOP 10", h2))
        kw_data = [["키워드", "빈도", "의미"]]
        for kw in result.get("keywords", [])[:10]:
            kw_data.append([
                str(kw.get("word", "")),
                str(kw.get("frequency", 0)),
                str(kw.get("meaning", "")),
            ])

        kw_table = Table(
            kw_data,
            colWidths=[35 * mm, 18 * mm, 122 * mm],
            repeatRows=1,
        )
        kw_table.setStyle(
            TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), body_font),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ])
        )
        story.append(kw_table)
        story.append(Spacer(1, 3 * mm))

        reasons = result.get("sentiment_reasons", {})
        for label, key in [
            ("긍정 주요 사유", "positive"),
            ("중립 주요 사유", "neutral"),
            ("부정·개선 요구 주요 사유", "negative"),
        ]:
            story.append(Paragraph(label, h2))
            vals = safe_list(reasons, key, 5)
            if vals:
                for v in vals:
                    story.append(Paragraph(f"• {v}", body))
            else:
                story.append(Paragraph("해당 의견 없음", body))

        story.append(Paragraph("개선 방안 5가지", h2))
        for i, item in enumerate(
            result.get("improvement_suggestions", [])[:5],
            start=1,
        ):
            story.append(Paragraph(f"{i}. {item}", body))

        story.append(Paragraph("개인별 피드백", h2))
        for fb in result.get("individual_feedback", [])[:10]:
            rid = fb.get("review_id", "")
            snippet = fb.get("review_snippet", "")
            feedback = fb.get("feedback", "")
            story.append(
                Paragraph(
                    f"• 교육생 의견 #{rid} - {snippet}",
                    body,
                )
            )
            story.append(
                Paragraph(
                    f"AI 피드백: {feedback}",
                    body,
                )
            )

        if idx < len(results):
            story.append(PageBreak())

    story.append(
        Spacer(1, 4 * mm)
    )
    story.append(
        Paragraph(
            "본 결과는 생성형 AI 기반 자동 분석 결과이며, 실제 업무 적용 시 담당자 최종 검토를 권장합니다.",
            body,
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


# =========================================================
# 4. 분석 프롬프트
# =========================================================
PROMPT_TEMPLATE = """
당신은 직업훈련기관의 강의평가·수강후기를 분석하는 데이터 분석가입니다.

아래는 하나의 교육과정에 대한 후기 목록입니다.
각 행에는 [리뷰ID], 선택적으로 [별점], [후기본문]이 있습니다.

[과정명]
{course_title}

[분석대상 후기 수]
{n_reviews}

[후기 목록]
{reviews_block}

다음 업무를 수행하십시오.

1. executive_summary
- 이 과정의 강점과 핵심 개선과제를 2~3문장으로 요약
- 후기에서 확인되는 내용만 사용

2. keywords
- 교육과정 개선에 의미 있는 핵심 키워드 정확히 10개
- word: 대표 키워드
- aliases: 후기에서 실제 함께 사용되는 유사 표현 0~5개
- meaning: 이 과정에서 해당 키워드가 의미하는 바 한 문장
- frequency는 반환하지 마십시오. Python이 실제 후기에서 다시 계산합니다.

3. review_sentiments
- 모든 {n_reviews}개 후기 각각을 positive / neutral / negative 중 하나로 분류
- review_id는 제공된 실제 [리뷰ID]를 그대로 사용
- 후기 수와 sentiment label 수가 정확히 일치해야 함

4. sentiment_reasons
- positive / neutral / negative의 대표적인 사유를 각각 최대 5개
- 해당 감성이 거의 없으면 빈 배열 허용
- 억지로 내용을 만들어내지 말 것

5. improvement_suggestions
- 후기 내용에 직접 근거한 구체적이고 실행 가능한 개선방안 정확히 5개

6. individual_feedback
- 리뷰가 10개 이상이면 반드시 서로 다른 리뷰 10개를 선택
- 리뷰가 10개 미만이면 가능한 리뷰 전부 선택
- review_id에는 제공된 실제 [리뷰ID] 사용
- review_snippet은 해당 리뷰 원문 일부(50자 이내)
- feedback은 해당 의견에 대응하는 정중하고 구체적인 피드백
- 단순 칭찬 반복보다 의견 반영, 추가 학습, 운영 보완 방향이 드러나게 작성

중요:
- 후기에서 확인되지 않는 사실을 만들지 마십시오.
- 개인정보를 추정하거나 생성하지 마십시오.
- 리뷰 ID와 후기 내용을 혼동하지 마십시오.
- 반드시 아래 JSON 구조만 반환하십시오.
- 코드블록, 머리말, 설명문을 붙이지 마십시오.

{{
  "executive_summary": "요약",
  "keywords": [
    {{
      "word": "키워드",
      "aliases": ["유사표현1", "유사표현2"],
      "meaning": "의미"
    }}
  ],
  "review_sentiments": [
    {{
      "review_id": "실제 리뷰ID",
      "sentiment": "positive"
    }}
  ],
  "sentiment_reasons": {{
    "positive": ["사유"],
    "neutral": ["사유"],
    "negative": ["사유"]
  }},
  "improvement_suggestions": [
    "개선방안"
  ],
  "individual_feedback": [
    {{
      "review_id": "실제 리뷰ID",
      "review_snippet": "원문 일부",
      "feedback": "개인별 피드백"
    }}
  ]
}}
"""


def analyze_course(group, course_title, review_col, rating_col, id_col):
    # 모든 유효 후기 사용
    sample = group.copy()

    if id_col:
        sample["_analysis_review_id"] = sample[id_col].astype(str)
    else:
        sample["_analysis_review_id"] = (sample.index + 1).astype(str)

    lines = []

    for _, row in sample.iterrows():
        review_id = row["_analysis_review_id"]

        rating_txt = ""
        if rating_col and pd.notna(row.get(rating_col)):
            rating_txt = f"[별점 {row[rating_col]}] "

        text = str(row[review_col]).replace("\n", " ").strip()

        lines.append(
            f"[리뷰ID {review_id}] {rating_txt}[후기본문] {text}"
        )

    prompt = PROMPT_TEMPLATE.format(
        course_title=course_title,
        n_reviews=len(sample),
        reviews_block="\n".join(lines),
    )

    response = client.responses.create(
        model=MODEL,
        input=prompt,
        max_output_tokens=14000,
    )

    result = extract_json(response.output_text)

    result.setdefault("executive_summary", "")
    result.setdefault("keywords", [])
    result.setdefault("review_sentiments", [])
    result.setdefault(
        "sentiment_reasons",
        {"positive": [], "neutral": [], "negative": []},
    )
    result.setdefault("improvement_suggestions", [])
    result.setdefault("individual_feedback", [])

    result["keywords"] = compute_keyword_frequencies(
        sample,
        review_col,
        result.get("keywords", [])[:10],
    )

    valid_ids = set(sample["_analysis_review_id"].astype(str))
    sentiment_map = {}

    for item in result.get("review_sentiments", []):
        rid = str(item.get("review_id", ""))
        sentiment = str(item.get("sentiment", "")).lower()

        if rid in valid_ids and sentiment in {
            "positive",
            "neutral",
            "negative",
        }:
            sentiment_map[rid] = sentiment

    counts = {
        "positive": sum(v == "positive" for v in sentiment_map.values()),
        "neutral": sum(v == "neutral" for v in sentiment_map.values()),
        "negative": sum(v == "negative" for v in sentiment_map.values()),
    }

    result["sentiment_counts"] = counts
    result["classified_reviews"] = len(sentiment_map)
    result["analysis_review_count"] = len(sample)

    return result


# =========================================================
# 5. 데이터 선택
# =========================================================
st.subheader("1. 분석 데이터")

source = st.radio(
    "데이터 사용 방법",
    ["샘플 데이터로 체험", "내 파일 업로드"],
    horizontal=True,
)

df = None
source_name = ""

if source == "샘플 데이터로 체험":
    if SAMPLE_FILE.exists():
        df = load_sample(str(SAMPLE_FILE))
        source_name = SAMPLE_FILE.name
        st.success(f"샘플 데이터 `{SAMPLE_FILE.name}`를 불러왔습니다.")
    else:
        st.warning(
            "`reviews_seoul_center_final.csv` 파일이 앱 폴더에 없습니다. "
            "GitHub 저장소에 app.py와 같은 위치로 올리거나 직접 업로드해주세요."
        )
        source = "내 파일 업로드"

if source == "내 파일 업로드":
    uploaded_file = st.file_uploader(
        "CSV 또는 Excel(.xlsx) 파일을 업로드하세요.",
        type=["csv", "xlsx"],
    )

    if uploaded_file is not None:
        try:
            df = load_dataframe(uploaded_file)
            source_name = uploaded_file.name
        except Exception as e:
            st.error(f"파일을 읽지 못했습니다: {e}")
            st.stop()

if df is None:
    st.info("샘플 파일을 저장소에 올리거나 분석할 파일을 업로드해주세요.")
    st.stop()


# =========================================================
# 6. 컬럼 인식
# =========================================================
course_col = find_col(
    df,
    exact_candidates=["title", "과정명", "훈련과정명", "course"],
    partial_candidates=["훈련과정"],
)

review_col = find_col(
    df,
    exact_candidates=[
        "content",
        "review_text",
        "reviewtext",
        "후기",
        "리뷰본문",
        "리뷰내용",
        "의견",
        "text",
    ],
    partial_candidates=["후기내용", "수강후기", "교육후기"],
    exclude=["review_id", "reviewid", "id"],
)

rating_col = find_col(
    df,
    exact_candidates=["rating", "별점", "평점", "score", "점수"],
)

id_col = find_col(
    df,
    exact_candidates=["review_id", "reviewid", "리뷰id", "후기id"],
)

c1, c2, c3, c4 = st.columns(4)
c1.metric("과정명 컬럼", course_col or "인식 실패")
c2.metric("후기본문 컬럼", review_col or "인식 실패")
c3.metric("별점 컬럼", rating_col or "없음")
c4.metric("리뷰ID 컬럼", id_col or "없음")

if course_col is None or review_col is None:
    st.error(
        "과정명 또는 후기본문 컬럼을 자동으로 찾지 못했습니다. "
        "예: title/과정명, content/후기 형태의 컬럼이 필요합니다."
    )
    st.stop()

df, removed_n = clean_dataframe(df, course_col, review_col, rating_col)

if df.empty:
    st.error("분석 가능한 후기 데이터가 없습니다.")
    st.stop()


# =========================================================
# 7. 전체 데이터 대시보드
# =========================================================
st.subheader("2. 데이터 요약")

overview = build_course_overview(df, course_col, rating_col)

m1, m2, m3, m4 = st.columns(4)
m1.metric("전체 후기", f"{len(df):,}건")
m2.metric("과정 수", f"{df[course_col].nunique():,}개")

if rating_col and df[rating_col].notna().any():
    m3.metric("전체 평균 별점", f"{df[rating_col].mean():.2f}")
else:
    m3.metric("전체 평균 별점", "-")

m4.metric("분석 모델", "GPT-5.6 Luna")

if removed_n > 0:
    st.caption(
        f"※ 과정명 또는 후기본문이 비어 있는 {removed_n:,}건은 분석에서 제외했습니다."
    )

left_overview, right_overview = st.columns([1.15, 0.85])

with left_overview:
    st.markdown("#### 과정별 현황")
    st.dataframe(
        overview,
        use_container_width=True,
        hide_index=True,
    )

with right_overview:
    st.markdown("#### 과정별 후기 수")
    course_count_chart = overview[["과정명", "후기수"]].set_index("과정명")
    st.bar_chart(course_count_chart, horizontal=True)

if rating_col and df[rating_col].notna().any():
    st.markdown("#### 별점 분포")

    rating_dist = (
        df[rating_col]
        .dropna()
        .round(1)
        .value_counts()
        .sort_index()
        .rename_axis("별점")
        .reset_index(name="후기수")
    )

    st.bar_chart(
        rating_dist.set_index("별점")["후기수"],
        use_container_width=True,
    )

with st.expander("원본 데이터 미리보기"):
    display_cols = [x for x in [id_col, course_col, rating_col, review_col] if x]
    st.dataframe(
        df[display_cols].head(30),
        use_container_width=True,
        hide_index=True,
    )


# =========================================================
# 8. 분석 실행
# =========================================================
st.subheader("3. AI 분석")

courses = df[course_col].dropna().astype(str).unique().tolist()

selected_courses = st.multiselect(
    "분석할 과정을 선택하세요.",
    options=courses,
    default=courses,
)

if st.button("🚀 AI 분석 시작", type="primary", use_container_width=True):
    if not selected_courses:
        st.warning("분석할 과정을 하나 이상 선택해주세요.")
        st.stop()

    results = {}
    errors = {}

    progress_bar = st.progress(0)
    status_box = st.empty()

    for i, course in enumerate(selected_courses, start=1):
        group = df[df[course_col].astype(str) == str(course)]
        status_box.info(
            f"🔄 {i}/{len(selected_courses)} 과정 분석 중 · {course} "
            f"({len(group):,}건 전체 분석)"
        )

        try:
            results[course] = analyze_course(
                group=group,
                course_title=course,
                review_col=review_col,
                rating_col=rating_col,
                id_col=id_col,
            )
        except Exception as e:
            errors[course] = str(e)

        progress_bar.progress(i / len(selected_courses))

    status_box.success(
        f"✅ 분석 완료 · 총 {len(selected_courses)}개 과정 처리"
    )
    time.sleep(0.2)

    st.session_state["results"] = results
    st.session_state["errors"] = errors
    st.session_state["source_name"] = source_name


# =========================================================
# 9. 결과 렌더링
# =========================================================
if "results" in st.session_state:
    st.subheader("4. 분석 결과")

    results = st.session_state["results"]
    errors = st.session_state.get("errors", {})

    if errors:
        with st.expander("⚠️ 분석 오류 확인"):
            for course, err in errors.items():
                st.error(f"{course}: {err}")

    for course, result in results.items():
        st.divider()
        st.header(f"📘 {course}")

        summary = result.get("executive_summary", "")
        if summary:
            st.info(f"**AI 핵심 요약**\n\n{summary}")

        counts = result.get(
            "sentiment_counts",
            {"positive": 0, "neutral": 0, "negative": 0},
        )

        pos = int(counts.get("positive", 0) or 0)
        neu = int(counts.get("neutral", 0) or 0)
        neg = int(counts.get("negative", 0) or 0)

        classified = int(result.get("classified_reviews", 0) or 0)
        analysis_n = int(result.get("analysis_review_count", 0) or 0)

        k1, k2, k3, k4 = st.columns(4)
        k1.metric("분석 후기", f"{analysis_n:,}건")

        if classified > 0:
            k2.metric("긍정", f"{pos:,}건", f"{pos / classified * 100:.1f}%")
            k3.metric("중립", f"{neu:,}건", f"{neu / classified * 100:.1f}%")
            k4.metric("부정", f"{neg:,}건", f"{neg / classified * 100:.1f}%")
        else:
            k2.metric("긍정", "-")
            k3.metric("중립", "-")
            k4.metric("부정", "-")

        if classified != analysis_n:
            st.caption(
                f"※ 감성 분류 응답 {classified:,}건 / 분석 대상 {analysis_n:,}건. "
                "모델 응답 누락이 있을 수 있으므로 최종 업무 적용 시 확인이 필요합니다."
            )

        left, right = st.columns([1.15, 0.85])

        with left:
            st.markdown("#### 🔑 핵심 키워드 TOP 10")

            keyword_rows = []

            for kw in result.get("keywords", [])[:10]:
                keyword_rows.append(
                    {
                        "키워드": kw.get("word", ""),
                        "빈도": int(kw.get("frequency", 0) or 0),
                        "의미": kw.get("meaning", ""),
                    }
                )

            if keyword_rows:
                keyword_df = pd.DataFrame(keyword_rows)

                st.dataframe(
                    keyword_df,
                    use_container_width=True,
                    hide_index=True,
                )

                st.bar_chart(
                    keyword_df.set_index("키워드")["빈도"],
                    horizontal=True,
                )
            else:
                st.info("키워드 분석 결과가 없습니다.")

        with right:
            st.markdown("#### 😊 감성 분포")

            sentiment_df = pd.DataFrame(
                {
                    "감성": ["긍정", "중립", "부정"],
                    "건수": [pos, neu, neg],
                }
            )

            st.bar_chart(
                sentiment_df.set_index("감성")["건수"],
                use_container_width=True,
            )

        reasons = result.get("sentiment_reasons", {})
        reason_cols = st.columns(3)

        with reason_cols[0]:
            st.markdown("##### 👍 긍정 주요 사유")
            values = safe_list(reasons, "positive", 5)
            if values:
                for item in values:
                    st.markdown(f"- {item}")
            else:
                st.caption("해당 의견 없음")

        with reason_cols[1]:
            st.markdown("##### ➖ 중립 주요 사유")
            values = safe_list(reasons, "neutral", 5)
            if values:
                for item in values:
                    st.markdown(f"- {item}")
            else:
                st.caption("해당 의견 없음")

        with reason_cols[2]:
            st.markdown("##### 👎 부정·개선 요구 주요 사유")
            values = safe_list(reasons, "negative", 5)
            if values:
                for item in values:
                    st.markdown(f"- {item}")
            else:
                st.caption("해당 의견 없음")

        st.markdown("#### 🛠️ 개선 방안 5가지")

        improvements = result.get("improvement_suggestions", [])[:5]
        if improvements:
            for idx, improvement in enumerate(improvements, start=1):
                st.markdown(f"**{idx}.** {improvement}")
        else:
            st.info("개선방안 결과가 없습니다.")

        st.markdown("#### 💬 개인별 피드백")

        feedbacks = result.get("individual_feedback", [])[:10]
        if feedbacks:
            for fb in feedbacks:
                review_id = fb.get("review_id", "")
                snippet = fb.get("review_snippet", "")
                title = (
                    f"교육생 의견 #{review_id} · {snippet}"
                    if review_id
                    else snippet
                )

                with st.expander(title or "개인별 피드백"):
                    st.markdown("**AI 피드백**")
                    st.write(fb.get("feedback", ""))
        else:
            st.info("개인별 피드백 결과가 없습니다.")

    st.divider()
    st.markdown("### 📄 분석 리포트 다운로드")

    download_cols = st.columns(3)

    with download_cols[0]:
        st.download_button(
            "📝 Word 리포트 다운로드",
            data=build_docx_report(results),
            file_name="수강후기_AI_분석_리포트.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    with download_cols[1]:
        try:
            pdf_bytes = build_pdf_report(results)
            st.download_button(
                "📕 PDF 리포트 다운로드",
                data=pdf_bytes,
                file_name="수강후기_AI_분석_리포트.pdf",
                mime="application/pdf",
                use_container_width=True,
            )
        except Exception as e:
            st.warning(f"PDF 생성 실패: {e}")

    with download_cols[2]:
        st.download_button(
            "🧾 JSON 원본 다운로드",
            data=json.dumps(results, ensure_ascii=False, indent=2),
            file_name="analysis_results.json",
            mime="application/json",
            use_container_width=True,
        )

    st.caption(
        "본 결과는 생성형 AI를 활용한 자동 분석 결과이며, "
        "실제 업무 적용 시 담당자의 최종 검토를 권장합니다."
    )
